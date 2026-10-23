import re
from pathlib import Path
from urllib.parse import unquote, urlparse

import duckdb
import pyarrow as pa
from openpyxl import load_workbook

from odf.services.datasource import get_dsn

SCHEMA = "file"
SUFFIXES = {".csv", ".tsv", ".txt", ".parquet", ".json", ".jsonl",
            ".ndjson", ".xlsx", ".xlsm"}


def _allowed_roots() -> list[Path]:
    from odf.settings import get_settings
    settings = get_settings()
    raw = settings.file_source_roots or settings.file_dir
    return [Path(item.strip()).expanduser().resolve()
            for item in raw.split(",") if item.strip()]


def resolve_path(dsn: str) -> Path:
    raw = (dsn or "").strip()
    if raw.startswith("file:"):
        parsed = urlparse(raw)
        raw = unquote(parsed.path)
        if re.match(r"^/[A-Za-z]:", raw):
            raw = raw[1:]
    if not raw:
        raise ValueError("file path is required")
    path = Path(raw).expanduser().resolve()
    roots = _allowed_roots()
    if not any(path == root or root in path.parents for root in roots):
        raise ValueError("file path is outside the configured file roots")
    if not path.is_file():
        raise ValueError("file not found: %s" % path)
    if path.suffix.lower() not in SUFFIXES:
        raise ValueError("unsupported file type: %s" % path.suffix)
    return path


def path_for(ds) -> Path:
    return resolve_path(get_dsn(ds))


def table_name_for(path: Path) -> str:
    name = re.sub(r"\W", "_", path.stem, flags=re.UNICODE)
    if not name or name[0].isdigit():
        name = "data_" + name
    return name[:100]


def _literal(value: str) -> str:
    return "'%s'" % value.replace("'", "''")


def _excel_table(path: Path) -> pa.Table:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header or all(v is None for v in header):
            raise ValueError("spreadsheet has no header row")
        names, seen = [], {}
        for i, value in enumerate(header, 1):
            name = str(value).strip() if value is not None else ""
            name = name or "column_%d" % i
            if name in seen:
                seen[name] += 1
                name = "%s_%d" % (name, seen[name])
            else:
                seen[name] = 1
            names.append(name)
        records = []
        for raw in rows:
            if raw is None or all(v is None for v in raw):
                continue
            records.append({name: raw[i] if i < len(raw) else None
                            for i, name in enumerate(names)})
        return pa.Table.from_pylist(records) if records else \
            pa.table({name: pa.array([], type=pa.null()) for name in names})
    finally:
        wb.close()


def _load(con, path: Path, name: str):
    suffix = path.suffix.lower()
    con.execute('CREATE SCHEMA IF NOT EXISTS %s' % SCHEMA)
    target = '"%s"."%s"' % (SCHEMA, name.replace('"', '""'))
    if suffix in (".xlsx", ".xlsm"):
        table = _excel_table(path)
        con.register("_odf_excel", table)
        try:
            con.execute("CREATE TABLE %s AS SELECT * FROM _odf_excel"
                        % target)
        finally:
            con.unregister("_odf_excel")
    elif suffix == ".parquet":
        con.execute("CREATE TABLE %s AS SELECT * FROM read_parquet(%s)"
                    % (target, _literal(str(path))))
    elif suffix in (".json", ".jsonl", ".ndjson"):
        con.execute("CREATE TABLE %s AS SELECT * FROM read_json_auto(%s)"
                    % (target, _literal(str(path))))
    else:
        con.execute("CREATE TABLE %s AS SELECT * FROM read_csv_auto(%s)"
                    % (target, _literal(str(path))))
    con.execute("SET search_path='%s'" % SCHEMA)
    con.execute("SET enable_external_access=false")
    con.execute("SET lock_configuration=true")


def connect(value) -> tuple[duckdb.DuckDBPyConnection, str]:
    path = value if isinstance(value, Path) else resolve_path(value)
    name = table_name_for(path)
    con = duckdb.connect()
    _load(con, path, name)
    return con, name


def describe(value) -> list[dict]:
    con = None
    try:
        con, name = connect(value)
        rows = con.execute('DESCRIBE SELECT * FROM "%s"."%s"'
                           % (SCHEMA, name)).fetchall()
        return [{"name": r[0], "type": r[1]} for r in rows]
    except ValueError:
        raise
    except Exception as e:
        raise ValueError("invalid file: %s" % e)
    finally:
        if con is not None:
            con.close()


def execute(ds, sql: str, params: dict | None = None,
            max_rows: int = 1000) -> dict:
    if not isinstance(sql, str) or not sql.strip():
        raise ValueError("sql is required")
    con = None
    binds = dict(params or {})
    values = []

    def repl(match):
        key = match.group(1)
        if key not in binds:
            raise ValueError("missing parameter: %s" % key)
        values.append(binds[key])
        return "?"

    try:
        con, _ = connect(get_dsn(ds))
        rendered = re.sub(r"(?<!:):([A-Za-z_]\w*)", repl, sql)
        result = con.execute(rendered, values)
        columns = [d[0] for d in (result.description or [])]
        rows = result.fetchmany(max_rows)
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(str(e))
    finally:
        if con is not None:
            con.close()
    return {"columns": columns,
            "rows": [dict(zip(columns, row)) for row in rows]}


def read_rows(ds, table: str, cursor_column: str | None, pk: str,
              cursor: str | None, limit: int):
    from odf.apistudio.compiler import quote_ident
    from odf.connectors.sql import format_cursor, parse_cursor
    path = path_for(ds)
    expected = table_name_for(path)
    name = (table or "").split(".")[-1]
    if name != expected:
        raise ValueError("table not found: %s" % table)
    con = None
    try:
        con, _ = connect(path)
        q = ("SELECT * FROM %s.%s" % (quote_ident(SCHEMA),
                                      quote_ident(name)))
        binds = []
        cur = parse_cursor(cursor)
        if cursor_column and cur:
            if len(cur) > 1 and cur[1] is not None:
                q += " WHERE (%s > ? OR (%s = ? AND %s > ?))" % (
                    quote_ident(cursor_column), quote_ident(cursor_column),
                    quote_ident(pk))
                binds += [cur[0], cur[0], cur[1]]
            else:
                q += " WHERE %s > ?" % quote_ident(cursor_column)
                binds.append(cur[0])
        if cursor_column:
            q += " ORDER BY %s, %s" % (quote_ident(cursor_column),
                                       quote_ident(pk))
        q += " LIMIT %d" % max(1, int(limit))
        result = con.execute(q, binds)
        columns = [d[0] for d in (result.description or [])]
        rows = [dict(zip(columns, r)) for r in result.fetchall()]
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(str(e))
    finally:
        if con is not None:
            con.close()
    new_cursor = cursor
    if cursor_column and rows:
        new_cursor = format_cursor(rows[-1][cursor_column], rows[-1][pk])
    return rows, new_cursor
